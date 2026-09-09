"""通过 COM 控制本机的 PowerPoint：打开文件、开始放映、翻页。

技术层：只有 COM 调用，不含任何演示业务概念。

**不缓存任何 COM 对象。** 调用方是 FastAPI 的线程池，同一份 ppt 的两次请求
可能落在不同线程上；COM 对象跨 apartment 使用需要 marshaling，win32com 不会
自动做，缓存下来的 Presentation 换个线程再用会直接报 RPC_E_WRONG_THREAD。
所以每个方法自己 CoInitialize + Dispatch —— Dispatch 会附着到已在运行的那个
PowerPoint 实例（没运行就启一个），状态全存在 PowerPoint 自己身上，Python
侧一个句柄都不留。代价是每次多几毫秒，换来彻底绕开线程亲和性。

    uv run python -m util.powerpoint_helper
"""

from __future__ import annotations

import sys
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parents[1]
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import time

import pythoncom
import win32api
import win32com.client
import win32con
import win32gui

from config.settings import PathConfig, SlideShowConfig, use_utf8_output


class PowerPointHelper:
    """一份 ppt 的放映控制。deck 只存文件路径，不存 COM 句柄。"""

    # 放映窗口的窗口类。PowerPoint 的编辑窗口是 PPTFrameClass，放映是这个。
    # SlideShowWindow 在部分 Office 版本上没有 HWND 属性（调用报 Member not
    # found），所以句柄从窗口类找，不走 COM。
    SHOW_WINDOW_CLASS = "screenClass"

    def __init__(self, **kwargs) -> None:
        self.deck = Path(kwargs.get("deck") or PathConfig.DECK).resolve()
        self.show_type = kwargs.get("show_type", SlideShowConfig.SPEAKER)
        self.advance_mode = kwargs.get("advance_mode", SlideShowConfig.MANUAL)
        # 放哪块屏。0 表示没指定，取主屏。
        self.monitor_index = kwargs.get("monitor") or SlideShowConfig.MONITOR

    # ---- COM 接入 ---------------------------------------------------------

    @staticmethod
    def _application():
        """当前线程的 PowerPoint 实例。没在运行的话这一行会把它启起来。"""
        pythoncom.CoInitialize()
        return win32com.client.Dispatch("PowerPoint.Application")

    def _presentation(self):
        """按文件名取，而不是按序号 —— 用户可能还开着别的 ppt。"""
        return self._application().Presentations(self.deck.name)

    def _view(self):
        return self._application().SlideShowWindows(1).View

    # ---- 状态 -------------------------------------------------------------

    @property
    def opened(self) -> bool:
        """这份 ppt 是否已经在 PowerPoint 里打开。"""
        application = self._application()
        return any(application.Presentations(index + 1).Name == self.deck.name
                   for index in range(application.Presentations.Count))

    @property
    def playing(self) -> bool:
        return self._application().SlideShowWindows.Count > 0

    @property
    def total(self) -> int:
        return self._presentation().Slides.Count if self.opened else 0

    @property
    def current(self) -> int:
        """当前页码。以 PowerPoint 的说法为准，自己不记一份影子计数。"""
        return self._view().CurrentShowPosition if self.playing else 0

    # ---- 窗口层级 ---------------------------------------------------------

    @classmethod
    def show_hwnd(cls) -> int:
        """放映窗口的句柄。"""
        return win32gui.FindWindow(cls.SHOW_WINDOW_CLASS, None)

    @property
    def foreground(self) -> bool:
        """放映窗口是不是当前前台窗口。"""
        return win32gui.GetForegroundWindow() == self.show_hwnd()

    def focus(self) -> None:
        """把放映窗口提到最前面。

        Run() 之后放映窗口并不一定在前台 —— 谁在前台就还是谁，放映就被盖住了。
        实测三种手段里只有 SetForegroundWindow 有效：SlideShowWindow.Activate()
        毫无反应，SetWindowPos(HWND_TOPMOST) 连 WS_EX_TOPMOST 都置不上。
        """
        handle = self.show_hwnd()
        win32gui.ShowWindow(handle, win32con.SW_SHOW)
        win32gui.SetForegroundWindow(handle)

    # ---- 显示器 -----------------------------------------------------------

    @staticmethod
    def monitors() -> list[tuple[int, int, int, int]]:
        """每块屏的像素矩形 (left, top, right, bottom)，编号就是这个列表的序号
        （从 1 开始）。副屏在主屏左边或上边时坐标是负的，属正常。"""
        return [tuple(rect) for _, _, rect in win32api.EnumDisplayMonitors()]

    @classmethod
    def primary_monitor(cls) -> int:
        """主屏的编号。Windows 里主屏的左上角固定是 (0, 0)。"""
        return next(index for index, rect in enumerate(cls.monitors(), start=1)
                    if rect[0] == 0 and rect[1] == 0)

    @property
    def monitor(self) -> int:
        """放映窗口现在在第几块屏。按窗口左上角落在谁的矩形里判断。"""
        if not self.playing:
            return 0
        window = self._application().SlideShowWindows(1)
        x = window.Left / SlideShowConfig.POINT_PER_PIXEL
        y = window.Top / SlideShowConfig.POINT_PER_PIXEL
        return next(index for index, (left, top, right, bottom)
                    in enumerate(self.monitors(), start=1)
                    if left <= x < right and top <= y < bottom)

    def move(self, monitor: int) -> int:
        """把放映窗口搬到第 monitor 块屏，铺满，仍是全屏放映。

        SlideShowWindow 的 Left/Top/Width/Height 可写，单位是磅，所以这里要把
        显示器的像素矩形换算过去。
        """
        left, top, right, bottom = self.monitors()[monitor - 1]
        scale = SlideShowConfig.POINT_PER_PIXEL
        window = self._application().SlideShowWindows(1)
        window.Left = left * scale
        window.Top = top * scale
        window.Width = (right - left) * scale
        window.Height = (bottom - top) * scale
        return monitor

    # ---- 打开与放映 -------------------------------------------------------

    def open(self) -> int:
        """打开 ppt 窗口，返回总页数。"""
        application = self._application()
        application.Visible = True
        return application.Presentations.Open(str(self.deck)).Slides.Count

    def run(self) -> int:
        """开始放映，返回当前页码。

        ShowPresenterView 必须先设：开着的话 PowerPoint 会自作主张把幻灯片扔到
        一块屏、备注面板扔到另一块，我们指定的屏上就未必是幻灯片了。
        """
        settings = self._presentation().SlideShowSettings
        settings.ShowType = self.show_type
        settings.AdvanceMode = self.advance_mode
        settings.ShowPresenterView = 0
        settings.Run()
        self.move(self.monitor_index or self.primary_monitor())
        self.focus()
        return self.current

    def exit(self) -> None:
        """退出放映，ppt 窗口还开着。"""
        self._view().Exit()

    def close(self) -> None:
        """关掉 ppt 文件。"""
        self._presentation().Close()

    # ---- 翻页 -------------------------------------------------------------

    def next(self) -> int:
        self._view().Next()
        return self.current

    def previous(self) -> int:
        self._view().Previous()
        return self.current

    def goto(self, slide: int) -> int:
        self._view().GotoSlide(slide)
        return self.current

    def screen(self, state: int) -> int:
        """黑屏 / 白屏 / 恢复正常，取 SlideShowConfig 里的 State 值。"""
        self._view().State = state
        return self.current


def demo_show(pause: float = 2.0, steps: int = 3) -> None:
    """放映配置里那份 ppt，翻几页，每步打印 PowerPoint 报回来的页码。

    只翻 steps 页就收，真实的 ppt 可能有几十页，验证 COM 通不通不需要走完。
    """
    deck = PathConfig.DECK
    if not deck.exists():
        make_demo_deck(deck)
        print(f"已生成示例 ppt: {deck}")

    powerpoint = PowerPointHelper(deck=deck)
    for index, rect in enumerate(powerpoint.monitors(), start=1):
        mark = "（主屏）" if index == powerpoint.primary_monitor() else ""
        print(f"屏幕 {index}: {rect}{mark}")

    print(f"\n打开 {powerpoint.deck.name} …")
    total = powerpoint.open()
    print(f"共 {total} 页")

    powerpoint.run()
    print(f"开始放映，在屏幕 {powerpoint.monitor} 上，当前第 {powerpoint.current} 页")
    print(f"放映窗口在最前面: {powerpoint.foreground}")

    for _ in range(min(steps, total - 1)):
        time.sleep(pause)
        print(f"下一页 -> 第 {powerpoint.next()} 页")

    time.sleep(pause)
    print(f"跳回第 {powerpoint.goto(1)} 页")

    time.sleep(pause)
    powerpoint.screen(SlideShowConfig.BLACK)
    print("黑屏")
    time.sleep(pause)
    powerpoint.screen(SlideShowConfig.RUNNING)
    print("恢复")

    time.sleep(pause)
    powerpoint.exit()
    print("退出放映")


def make_demo_deck(path: Path) -> Path:
    """给 demo 用的三页 ppt。真实使用时把自己的文件路径写进 .env。"""
    from pptx import Presentation

    presentation = Presentation()
    layout = presentation.slide_layouts[1]
    for index, title in enumerate(["第一页", "第二页", "第三页"], start=1):
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = f"这是第 {index} 页，用来验证翻页控制。"

    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(path))
    return path


def main() -> None:
    use_utf8_output()
    demo_show()


if __name__ == "__main__":
    main()
